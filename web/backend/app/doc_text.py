"""Extract plain text from PDF / Excel / Word for cheap OpenAI prompts."""

from __future__ import annotations

from pathlib import Path


def extract_file_text(path: Path, max_chars: int = 20000) -> str:
    if not path.is_file():
        return ""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            text = _pdf_text(path)
        elif ext in {".xlsx", ".xlsm", ".xls"}:
            text = _excel_text(path)
        elif ext in {".docx", ".doc"}:
            text = _word_text(path)
        elif ext in {".txt", ".csv", ".rtf"}:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            return ""
    except Exception as exc:
        return f"[Could not read {path.name}: {exc}]"
    text = (text or "").strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n…[truncated for cost control]"
    return text


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages[:40]:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _excel_text(path: Path) -> str:
    from openpyxl import load_workbook

    # data_only reads cached values when present
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets[:8]:
        lines.append(f"=== Sheet: {sheet.title} ===")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                lines.append("\t".join(cells))
                row_count += 1
            if row_count >= 200:
                lines.append("…[sheet rows truncated]")
                break
    wb.close()
    return "\n".join(lines)


def _word_text(path: Path) -> str:
    if path.suffix.lower() == ".doc":
        # Legacy .doc — best-effort binary strip
        raw = path.read_bytes()
        return raw.decode("latin-1", errors="ignore")
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables[:20]:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)
